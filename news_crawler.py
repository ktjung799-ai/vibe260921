"""네이버 뉴스 크롤러 (PyQt6 GUI)

검색어(또는 네이버 검색 URL)를 입력하면 검색 결과의 뉴스 기사를 크롤링해
제목/언론사/날짜/본문을 보여주고 JSON, CSV, 엑셀(xlsx)로 저장할 수 있다.

    pip install requests beautifulsoup4 PyQt6 openpyxl
    python news_crawler.py
"""
import csv
import json
import sys
import time
from urllib.parse import quote

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication, QFileDialog, QHBoxLayout, QHeaderView, QLabel, QLineEdit,
    QMainWindow, QMessageBox, QProgressBar, QPushButton, QSpinBox, QSplitter,
    QTableWidget, QTableWidgetItem, QTextEdit, QVBoxLayout, QWidget,
    QAbstractItemView,
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
}
NEWS_HOST = "n.news.naver.com/mnews/article/"
DELAY = 1.0  # 요청 사이 대기(초): 서버에 부담을 주지 않도록


# ---------------------------------------------------------------- 크롤링 로직
def fetch(session, url):
    resp = session.get(url, headers=HEADERS, timeout=10)
    resp.raise_for_status()
    resp.encoding = "utf-8"
    return BeautifulSoup(resp.text, "html.parser")


def get_article_links(soup, limit):
    """검색 결과에서 네이버 뉴스 기사 링크를 중복 없이 추출한다.

    클래스명이 난독화되어 자주 바뀌므로 링크 주소 패턴으로 찾는다.
    """
    links = []
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if NEWS_HOST in href:
            href = href.split("?")[0]  # sid 등 쿼리 제거
            if href not in links:
                links.append(href)
        if len(links) >= limit:
            break
    return links


def text_of(soup, selector):
    el = soup.select_one(selector)
    return el.get_text(strip=True) if el else ""


def parse_article(soup, url):
    body_el = soup.select_one("#dic_area")
    if body_el:
        # 본문 안의 사진 설명, 스크립트 등 불필요한 요소 제거
        for tag in body_el.select("script, style, .img_desc, .end_photo_org, .media_end_summary"):
            tag.decompose()
        body = body_el.get_text("\n", strip=True)
    else:
        body = ""

    date_el = soup.select_one("span.media_end_head_info_datestamp_time")
    press = soup.select_one("meta[property='og:article:author']")

    return {
        "url": url,
        "title": text_of(soup, "#title_area"),
        "press": press["content"].split("|")[0].strip() if press else "",
        "date": date_el.get("data-date-time", date_el.get_text(strip=True)) if date_el else "",
        "body": body,
    }


def save_xlsx(path, articles):
    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"
    ws.append(["번호", "제목", "언론사", "날짜", "URL", "본문"])
    for i, a in enumerate(articles, 1):
        ws.append([i, a["title"], a["press"], a["date"], a["url"], a["body"][:32000]])  # 셀 글자 수 제한(32767)
        ws.cell(row=i + 1, column=5).hyperlink = a["url"]
    head_fill = PatternFill("solid", fgColor="4F46E5")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for col, width in zip("ABCDEF", (6, 50, 14, 20, 40, 80)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (2, 6))
        ws.row_dimensions[row[0].row].height = 60
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def build_search_url(text):
    text = text.strip()
    if text.startswith(("http://", "https://")):
        return text
    return f"https://search.naver.com/search.naver?where=news&query={quote(text)}"


# ---------------------------------------------------------------- 백그라운드 작업
class CrawlWorker(QThread):
    log = pyqtSignal(str)
    article = pyqtSignal(dict)
    progress = pyqtSignal(int, int)  # 현재, 전체
    failed = pyqtSignal(str)

    def __init__(self, url, limit):
        super().__init__()
        self.url = url
        self.limit = limit
        self._stop = False

    def stop(self):
        self._stop = True

    def run(self):
        try:
            session = requests.Session()
            self.log.emit("검색 결과를 불러오는 중...")
            links = get_article_links(fetch(session, self.url), self.limit)
        except requests.RequestException as e:
            self.failed.emit(f"검색 페이지를 불러오지 못했습니다.\n{e}")
            return
        if not links:
            self.failed.emit("네이버 뉴스 기사 링크를 찾지 못했습니다.\n(검색 결과가 없거나 페이지 구조 변경/접근 제한일 수 있습니다)")
            return

        total = len(links)
        self.log.emit(f"기사 {total}건을 찾았습니다.")
        self.progress.emit(0, total)
        for i, link in enumerate(links, 1):
            if self._stop:
                self.log.emit("사용자가 중지했습니다.")
                return
            try:
                time.sleep(DELAY)
                art = parse_article(fetch(session, link), link)
            except requests.RequestException as e:
                self.log.emit(f"[{i}/{total}] 실패: {link} ({e})")
            else:
                if art["body"]:
                    self.article.emit(art)
                    self.log.emit(f"[{i}/{total}] {art['title']}")
                else:
                    self.log.emit(f"[{i}/{total}] 본문 없음(건너뜀): {link}")
            self.progress.emit(i, total)
        self.log.emit("크롤링 완료")


# ---------------------------------------------------------------- GUI
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("네이버 뉴스 크롤러")
        self.resize(1000, 700)
        self.articles = []
        self.worker = None

        # 입력 영역
        self.input = QLineEdit("반도체")
        self.input.setPlaceholderText("검색어 또는 네이버 검색 URL")
        self.input.returnPressed.connect(self.start)
        self.spin = QSpinBox()
        self.spin.setRange(1, 30)
        self.spin.setValue(5)
        self.spin.setSuffix(" 건")
        self.start_btn = QPushButton("크롤링 시작")
        self.start_btn.clicked.connect(self.start)
        self.stop_btn = QPushButton("중지")
        self.stop_btn.setEnabled(False)
        self.stop_btn.clicked.connect(self.stop)

        top = QHBoxLayout()
        top.addWidget(QLabel("검색어/URL"))
        top.addWidget(self.input, 1)
        top.addWidget(self.spin)
        top.addWidget(self.start_btn)
        top.addWidget(self.stop_btn)

        # 결과 목록 + 본문
        self.table = QTableWidget(0, 3)
        self.table.setHorizontalHeaderLabels(["제목", "언론사", "날짜"])
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.verticalHeader().setVisible(False)
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        self.table.itemSelectionChanged.connect(self.show_detail)

        self.detail = QTextEdit()
        self.detail.setReadOnly(True)
        self.detail.setPlaceholderText("목록에서 기사를 선택하면 본문이 표시됩니다.")

        split = QSplitter(Qt.Orientation.Vertical)
        split.addWidget(self.table)
        split.addWidget(self.detail)
        split.setStretchFactor(0, 2)
        split.setStretchFactor(1, 3)

        # 로그
        self.log_box = QTextEdit()
        self.log_box.setReadOnly(True)
        self.log_box.setMaximumHeight(90)

        # 저장 버튼 + 진행 상태
        self.json_btn = QPushButton("JSON 저장")
        self.csv_btn = QPushButton("CSV 저장")
        self.xlsx_btn = QPushButton("엑셀 저장")
        self.json_btn.clicked.connect(lambda: self.save("json"))
        self.csv_btn.clicked.connect(lambda: self.save("csv"))
        self.xlsx_btn.clicked.connect(lambda: self.save("xlsx"))
        self.progress = QProgressBar()
        self.progress.setValue(0)
        bottom = QHBoxLayout()
        bottom.addWidget(self.progress, 1)
        bottom.addWidget(self.json_btn)
        bottom.addWidget(self.csv_btn)
        bottom.addWidget(self.xlsx_btn)
        self.update_save_buttons()

        layout = QVBoxLayout()
        layout.addLayout(top)
        layout.addWidget(split, 1)
        layout.addWidget(self.log_box)
        layout.addLayout(bottom)
        central = QWidget()
        central.setLayout(layout)
        self.setCentralWidget(central)

    # ---- 동작
    def update_save_buttons(self):
        has = bool(self.articles)
        self.json_btn.setEnabled(has)
        self.csv_btn.setEnabled(has)
        self.xlsx_btn.setEnabled(has)

    def set_running(self, running):
        self.start_btn.setEnabled(not running)
        self.stop_btn.setEnabled(running)
        self.input.setEnabled(not running)
        self.spin.setEnabled(not running)

    def log(self, msg):
        self.log_box.append(msg)

    def start(self):
        text = self.input.text().strip()
        if not text:
            QMessageBox.warning(self, "입력 필요", "검색어 또는 URL을 입력해 주세요.")
            return
        if self.worker and self.worker.isRunning():
            return
        self.articles.clear()
        self.table.setRowCount(0)
        self.detail.clear()
        self.log_box.clear()
        self.progress.setValue(0)
        self.update_save_buttons()

        self.worker = CrawlWorker(build_search_url(text), self.spin.value())
        self.worker.log.connect(self.log)
        self.worker.article.connect(self.add_article)
        self.worker.progress.connect(self.on_progress)
        self.worker.failed.connect(self.on_failed)
        self.worker.finished.connect(self.on_finished)
        self.set_running(True)
        self.worker.start()

    def stop(self):
        if self.worker:
            self.worker.stop()
            self.stop_btn.setEnabled(False)

    def add_article(self, art):
        self.articles.append(art)
        row = self.table.rowCount()
        self.table.insertRow(row)
        for col, key in enumerate(("title", "press", "date")):
            self.table.setItem(row, col, QTableWidgetItem(art[key]))
        self.update_save_buttons()

    def on_progress(self, cur, total):
        self.progress.setMaximum(total)
        self.progress.setValue(cur)

    def on_failed(self, msg):
        self.log(msg.replace("\n", " "))
        QMessageBox.warning(self, "크롤링 실패", msg)

    def on_finished(self):
        self.set_running(False)

    def show_detail(self):
        rows = self.table.selectionModel().selectedRows()
        if not rows:
            return
        art = self.articles[rows[0].row()]
        self.detail.setPlainText(
            f"{art['title']}\n\n{art['press']}  |  {art['date']}\n{art['url']}\n\n{art['body']}"
        )

    def save(self, kind):
        if not self.articles:
            return
        filt = {"json": "JSON (*.json)", "csv": "CSV (*.csv)", "xlsx": "Excel (*.xlsx)"}[kind]
        path, _ = QFileDialog.getSaveFileName(self, "저장", f"news_result.{kind}", filt)
        if not path:
            return
        try:
            if kind == "json":
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(self.articles, f, ensure_ascii=False, indent=2)
            elif kind == "xlsx":
                save_xlsx(path, self.articles)
            else:
                with open(path, "w", encoding="utf-8-sig", newline="") as f:  # 엑셀 한글 호환
                    w = csv.DictWriter(f, fieldnames=["title", "press", "date", "url", "body"])
                    w.writeheader()
                    w.writerows(self.articles)
        except (OSError, ValueError) as e:  # 파일이 엑셀에서 열려 있는 경우 등
            QMessageBox.critical(self, "저장 실패", str(e))
            return
        self.log(f"저장 완료: {path}")

    def closeEvent(self, event):
        if self.worker and self.worker.isRunning():
            self.worker.stop()
            self.worker.wait(3000)
        event.accept()


def main():
    app = QApplication(sys.argv)
    win = MainWindow()
    win.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
