"""코스피200 데이터 크롤러 (네이버 증권 stock.naver.com/market/stock/kr 기준)

대상 페이지: https://stock.naver.com/market/stock/kr

[BeautifulSoup4를 쓰지 않는 이유]
이 페이지는 자바스크립트가 화면을 그리는 구조라 받아온 HTML에는 종목 데이터도, '코스피200'이라는
글자도 없다. BeautifulSoup4는 HTML을 파싱하는 도구라 이 데이터를 가져올 수 없다.
그래서 페이지가 내부적으로 호출하는 JSON API를 requests로 직접 호출한다.

[수집 방법]
이 페이지의 종목 목록은 시장 구분이 전체/코스피/코스닥뿐이고 '코스피200' 필터가 없다.
1) 코스피200 편입종목 코드         : m.stock.naver.com/api/index/KPI200/enrollStocks
2) 코스피 전체 종목의 상세 시세    : stock.naver.com/api/domestic/market/stock/default
두 결과를 종목코드로 합쳐, 코스피200 종목의 상세 데이터(PER, PBR, ROE, 52주 고저 등)를 만든다.
3) 코스피200 지수 자체의 시세       : m.stock.naver.com/api/index/KPI200/basic

사용법:
    pip install requests openpyxl
    python kospi200_market.py                 # 시가총액 상위 20개 출력, CSV 저장
    python kospi200_market.py -n 50           # 상위 50개 출력
    python kospi200_market.py --xlsx          # 엑셀로도 저장
CSV/엑셀에는 편입종목 전체가 저장된다. (지수 조정 시기에 따라 200개가 아닌 201개가 내려올 수 있다)
"""
import argparse
import csv
import sys
import time

import requests

INDEX_CODE = "KPI200"
URL_CONSTITUENTS = "https://m.stock.naver.com/api/index/{code}/enrollStocks"
URL_INDEX = "https://m.stock.naver.com/api/index/{code}/basic"
URL_MARKET = "https://stock.naver.com/api/domestic/market/stock/default"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Referer": "https://stock.naver.com/market/stock/kr",
    "Accept": "application/json",
}
PAGE_SIZE = 50   # 편입종목 API가 한 번에 허용하는 최대치
DELAY = 0.5      # 요청 사이 대기(초)

# (컬럼명, API 필드, 변환)   변환: 원 단위를 억/백만 단위로 바꾸는 등
COLUMNS = [
    ("순위", None, None),
    ("종목코드", "itemcode", None),
    ("종목명", "itemname", None),
    ("현재가", "nowPrice", None),
    ("전일비", "prevChangePrice", None),
    ("등락률(%)", "prevChangeRate", None),
    ("시가", "openPrice", None),
    ("고가", "highPrice", None),
    ("저가", "lowPrice", None),
    ("거래량", "tradeVolume", None),
    ("거래대금(백만)", "tradeAmount", 1_000_000),
    ("시가총액(억)", "marketSum", 100_000_000),
    ("외국인보유율(%)", "frgnHoldRate", None),
    ("PER", "per", None),
    ("PBR", "pbr", None),
    ("ROE(%)", "roe", None),
    ("EPS", "eps", None),
    ("배당수익률(%)", "dividendRate", None),
    ("52주최고", "week52HighPrice", None),
    ("52주최저", "week52LowPrice", None),
]
FIELDS = [c[0] for c in COLUMNS]


def get_json(session, url, params=None):
    resp = session.get(url, params=params, headers=HEADERS, timeout=15)
    resp.raise_for_status()
    return resp.json()


def to_num(text, divisor=None):
    """'272,500' / '4.6' -> 숫자, 변환 불가(None, 'N/A' 등)이면 None"""
    try:
        value = float(str(text).replace(",", ""))
    except (TypeError, ValueError):
        return None
    if divisor:
        return round(value / divisor, 2)
    return int(value) if value.is_integer() else value


def fetch_constituent_codes(session):
    """코스피200 편입종목 코드를 시가총액 순서대로 가져온다."""
    codes, page = [], 1
    while True:
        items = get_json(session, URL_CONSTITUENTS.format(code=INDEX_CODE),
                         {"page": page, "pageSize": PAGE_SIZE})
        if not items:
            break
        codes.extend(i["itemCode"] for i in items)
        if len(items) < PAGE_SIZE:
            break
        page += 1
        time.sleep(DELAY)
    return codes


def fetch_kospi_market(session):
    """코스피 전체 종목의 상세 시세를 시가총액 순서로 가져온다 (한 번에 최대 1000건)."""
    return get_json(session, URL_MARKET, {
        "tradeType": "KRX", "marketType": "KOSPI", "orderType": "marketSum",
        "startIdx": 0, "pageSize": 1000,
    })


def fetch_index_summary(session):
    d = get_json(session, URL_INDEX.format(code=INDEX_CODE))
    return {
        "name": d.get("stockName", "코스피 200"),
        "price": d.get("closePrice"),
        "diff": d.get("compareToPreviousClosePrice"),
        "ratio": d.get("fluctuationsRatio"),
        "time": d.get("localTradedAt"),
    }


def build_rows(codes, market):
    """편입종목 코드와 코스피 시세를 합쳐 코스피200 종목만 시가총액순으로 만든다."""
    by_code = {m["itemcode"]: m for m in market}
    rows, missing = [], []
    for code in codes:
        m = by_code.get(code)
        if m is None:
            missing.append(code)
            continue
        row = {}
        for name, key, div in COLUMNS:
            row[name] = None if key is None else to_num(m.get(key), div)
        row["종목코드"], row["종목명"] = m["itemcode"], m["itemname"]
        rows.append(row)
    rows.sort(key=lambda r: r["시가총액(억)"] or 0, reverse=True)
    for i, r in enumerate(rows, 1):
        r["순위"] = i
    return rows, missing


def crawl():
    session = requests.Session()
    codes = fetch_constituent_codes(session)
    time.sleep(DELAY)
    market = fetch_kospi_market(session)
    time.sleep(DELAY)
    summary = fetch_index_summary(session)
    rows, missing = build_rows(codes, market)
    return summary, rows, missing


def fmt(v, digits=0):
    return "-" if v is None else f"{v:,.{digits}f}"


def print_table(rows):
    print(f"{'순위':>4} {'종목명':<14}{'현재가':>10}{'전일비':>9}{'등락률':>8}"
          f"{'시가총액(억)':>14}{'PER':>8}{'PBR':>7}{'ROE':>7}")
    for r in rows:
        print(f"{r['순위']:>4} {r['종목명']:<14.14}{fmt(r['현재가']):>10}{fmt(r['전일비']):>9}"
              f"{fmt(r['등락률(%)'], 2):>7}%{fmt(r['시가총액(억)']):>14}"
              f"{fmt(r['PER'], 2):>8}{fmt(r['PBR'], 2):>7}{fmt(r['ROE(%)'], 2):>7}")


def save_csv(rows, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:  # 엑셀 한글 호환
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)


def save_xlsx(rows, summary, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "KOSPI200"
    ws.append(FIELDS)
    for r in rows:
        ws.append([r[k] for k in FIELDS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F46E5")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for col in range(4, len(FIELDS) + 1):
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = "#,##0.00" if FIELDS[col - 1] in (
                    "등락률(%)", "외국인보유율(%)", "PER", "PBR", "ROE(%)", "배당수익률(%)") else "#,##0"
    ws.column_dimensions["C"].width = 22
    for idx in range(4, len(FIELDS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = 14
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("지수")
    info.append(["항목", "값"])
    for k, v in (("지수", summary["name"]), ("현재값", summary["price"]),
                 ("전일대비", summary["diff"]), ("등락률(%)", summary["ratio"]),
                 ("기준시각", summary["time"])):
        info.append([k, v])
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 28
    wb.save(path)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")  # Windows 콘솔 한글 깨짐 방지

    parser = argparse.ArgumentParser(description="코스피200 데이터 크롤러")
    parser.add_argument("-n", "--num", type=int, default=20, help="화면에 출력할 종목 수 (기본 20)")
    parser.add_argument("--xlsx", action="store_true", help="엑셀 파일로도 저장")
    args = parser.parse_args()

    try:
        summary, rows, missing = crawl()
    except requests.RequestException as e:
        sys.exit(f"데이터를 가져오지 못했습니다: {e}")
    except (ValueError, KeyError, TypeError):
        sys.exit("응답을 해석하지 못했습니다. (API 주소나 응답 구조가 바뀌었을 수 있습니다)")
    if not rows:
        sys.exit("가져온 종목이 없습니다.")

    print(f"[{summary['name']}] {summary['price']}  전일대비 {summary['diff']} ({summary['ratio']}%)  "
          f"기준 {summary['time']}\n")
    print_table(rows[:args.num])
    if missing:
        print(f"\n※ 코스피 시세에서 찾지 못한 편입종목 {len(missing)}개 제외: {', '.join(missing)}")

    save_csv(rows, "kospi200_market.csv")
    print(f"\n{len(rows)}개 종목 저장 완료: kospi200_market.csv")
    if args.xlsx:
        save_xlsx(rows, summary, "kospi200_market.xlsx")
        print("엑셀 저장 완료: kospi200_market.xlsx (시트: KOSPI200, 지수)")


if __name__ == "__main__":
    main()
